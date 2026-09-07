"""
Guia "Resultados" dentro do mesmo simulado.xlsx (ao lado das guias de eixo e
do Consolidado). Cada linha = uma resposta de um participante a UMA questão.

Colunas: timestamp, simulacao, participante, eixo, tema, pergunta,
         resposta_dada (texto da alternativa escolhida),
         resposta_correta (texto da alternativa certa),
         acertou (True/False)

Isso permite cruzar desempenho por: participante, eixo, tema, simulado
específico ou evolução ao longo de vários simulados — sem precisar de
Google Forms.
"""
import datetime as dt

from openpyxl import Workbook

CABECALHO = [
    "timestamp", "simulacao", "participante", "eixo", "tema",
    "pergunta", "resposta_dada", "resposta_correta", "acertou",
]


def _guia_resultados(wb: Workbook):
    if "Resultados" not in wb.sheetnames:
        ws = wb.create_sheet("Resultados")
        ws.append(CABECALHO)
    return wb["Resultados"]


def listar_simulacoes_disponiveis(wb: Workbook) -> list:
    """Números de simulação presentes no Consolidado, do mais recente pro mais antigo."""
    if "Consolidado" not in wb.sheetnames:
        return []
    ws = wb["Consolidado"]
    numeros = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            numeros.add(int(row[0]))
    return sorted(numeros, reverse=True)


def carregar_questoes_da_simulacao(wb: Workbook, numero_simulacao: int) -> list:
    """
    Lê o Consolidado e devolve as questões de uma simulação específica, no
    formato: [{"eixo", "tema", "pergunta", "opcoes": [...], "correta": int}, ...]
    """
    ws = wb["Consolidado"]
    questoes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        simulacao, eixo, pergunta = row[0], row[1], row[2]
        if simulacao != numero_simulacao:
            continue
        opcoes = [row[3], row[4], row[5], row[6], row[7]]
        correta_letra = row[8]
        tema = row[10] if len(row) > 10 else ""

        # "correta" no Consolidado é gravada como a LETRA (A-E); convertemos
        # para índice 0-4 pra facilitar a correção depois.
        letras = ["A", "B", "C", "D", "E"]
        if isinstance(correta_letra, str) and correta_letra.strip().upper() in letras:
            idx_correta = letras.index(correta_letra.strip().upper())
        else:
            idx_correta = int(correta_letra) if correta_letra is not None else 0

        questoes.append({
            "eixo": eixo,
            "tema": tema,
            "pergunta": pergunta,
            "opcoes": opcoes,
            "correta": idx_correta,
        })
    return questoes


def gravar_respostas(wb: Workbook, numero_simulacao: int, participante: str, respostas: list):
    """
    respostas: [{"eixo", "tema", "pergunta", "opcoes", "correta", "resposta_idx"}, ...]
    (o mesmo formato de carregar_questoes_da_simulacao, com "resposta_idx" a mais)
    """
    ws = _guia_resultados(wb)
    agora = dt.datetime.now().isoformat(timespec="seconds")

    for r in respostas:
        resposta_idx = r.get("resposta_idx")
        resposta_texto = r["opcoes"][resposta_idx] if resposta_idx is not None else ""
        correta_texto = r["opcoes"][r["correta"]]
        acertou = resposta_idx == r["correta"]

        ws.append([
            agora, numero_simulacao, participante, r["eixo"], r["tema"],
            r["pergunta"], resposta_texto, correta_texto, acertou,
        ])


def carregar_todos_resultados(wb: Workbook) -> list:
    """Devolve todas as linhas da guia Resultados como lista de dicts."""
    if "Resultados" not in wb.sheetnames:
        return []
    ws = wb["Resultados"]
    linhas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        linhas.append(dict(zip(CABECALHO, row)))
    return linhas
