"""
Camada de armazenamento no Google Drive — equivalente ao que config_manager.py
e excel_manager.py faziam com Path/disco local, agora usando o token de longa
duração da sua conta (drive_service.py) em vez de acesso a disco.

Estrutura no Drive (dentro da pasta raiz, em st.secrets["google_drive"]["root_folder_id"]):
  <pasta raiz>/
    participantes.json        (lista de quem pode responder + quais estudos)
    <slug_do_estudo_1>/
        config.json
        simulado.xlsx          (guias por eixo + Consolidado + Resultados)
    <slug_do_estudo_2>/
        ...
"""
import io
import json

import streamlit as st
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from openpyxl import Workbook, load_workbook

import drive_service as dsvc

MIME_FOLDER = "application/vnd.google-apps.folder"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def root_folder_id() -> str:
    return st.secrets["google_drive"]["root_folder_id"]


def _buscar(nome: str, parent_id: str, mime: str | None = None) -> dict | None:
    service = dsvc.obter_service()
    nome_escapado = nome.replace("'", r"\'")
    q = f"name = '{nome_escapado}' and '{parent_id}' in parents and trashed = false"
    if mime:
        q += f" and mimeType = '{mime}'"
    resp = service.files().list(q=q, fields="files(id, name, mimeType)").execute()
    arquivos = resp.get("files", [])
    return arquivos[0] if arquivos else None


# Wrappers públicos de _buscar/_baixar_bytes/_upload_bytes — reutilizados por
# outros módulos (ex: participantes_manager.py) que também precisam ler/
# gravar um arquivo solto na pasta raiz, sem duplicar a lógica de busca.
def buscar_arquivo(nome: str, parent_id: str, mime: str | None = None) -> dict | None:
    return _buscar(nome, parent_id, mime)


def baixar_bytes(file_id: str) -> bytes:
    return _baixar_bytes(file_id)


def salvar_bytes(nome: str, dados: bytes, parent_id: str, mime_type: str,
                  file_id: str | None = None) -> dict:
    return _upload_bytes(nome, dados, parent_id, mime_type, file_id=file_id)


def listar_estudos() -> list:
    """Lista as subpastas (cada uma = um estudo) dentro da pasta raiz."""
    return listar_subpastas(root_folder_id())


def listar_subpastas(parent_id: str) -> list:
    """Versão genérica de listar_estudos, reutilizável para qualquer pasta."""
    service = dsvc.obter_service()
    resp = service.files().list(
        q=f"'{parent_id}' in parents and mimeType = '{MIME_FOLDER}' and trashed = false",
        fields="files(id, name)",
    ).execute()
    return sorted(resp.get("files", []), key=lambda f: f["name"])


def listar_arquivos(parent_id: str, apenas_extensao: str | None = None) -> list:
    """Lista arquivos (não-pasta) soltos dentro de uma pasta."""
    service = dsvc.obter_service()
    q = f"'{parent_id}' in parents and mimeType != '{MIME_FOLDER}' and trashed = false"
    resp = service.files().list(q=q, fields="files(id, name)").execute()
    arquivos = resp.get("files", [])
    if apenas_extensao:
        arquivos = [a for a in arquivos if a["name"].endswith(apenas_extensao)]
    return arquivos


def obter_ou_criar_subpasta(nome: str, parent_id: str) -> str:
    """Versão genérica: acha ou cria uma subpasta com esse nome dentro de QUALQUER pasta pai."""
    existente = _buscar(nome, parent_id, MIME_FOLDER)
    if existente:
        return existente["id"]

    service = dsvc.obter_service()
    metadata = {"name": nome, "mimeType": MIME_FOLDER, "parents": [parent_id]}
    criado = service.files().create(body=metadata, fields="id").execute()
    return criado["id"]


def obter_ou_criar_pasta_estudo(slug: str) -> str:
    return obter_ou_criar_subpasta(slug, root_folder_id())


def mover_para_lixeira(file_id: str):
    """Move um arquivo OU pasta (com tudo dentro) para a lixeira do Drive — não é exclusão permanente."""
    service = dsvc.obter_service()
    service.files().update(fileId=file_id, body={"trashed": True}).execute()


def _baixar_bytes(file_id: str) -> bytes:
    service = dsvc.obter_service()
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, service.files().get_media(fileId=file_id))
    concluido = False
    while not concluido:
        _, concluido = downloader.next_chunk()
    return buffer.getvalue()


def _upload_bytes(nome: str, dados: bytes, parent_id: str, mime_type: str,
                   file_id: str | None = None) -> dict:
    service = dsvc.obter_service()
    media = MediaIoBaseUpload(io.BytesIO(dados), mimetype=mime_type, resumable=False)

    if file_id:
        return service.files().update(fileId=file_id, media_body=media).execute()

    metadata = {"name": nome, "parents": [parent_id]}
    return service.files().create(body=metadata, media_body=media, fields="id").execute()


# ---------------- config.json ----------------

def ler_config(pasta_estudo_id: str) -> dict:
    arquivo = _buscar("config.json", pasta_estudo_id)
    if not arquivo:
        raise FileNotFoundError("config.json não encontrado na pasta do estudo no Drive.")
    dados = _baixar_bytes(arquivo["id"])
    return json.loads(dados.decode("utf-8"))


def salvar_config(pasta_estudo_id: str, config: dict):
    arquivo = _buscar("config.json", pasta_estudo_id)
    dados = json.dumps(config, ensure_ascii=False, indent=2).encode("utf-8")
    _upload_bytes("config.json", dados, pasta_estudo_id, "application/json",
                  file_id=arquivo["id"] if arquivo else None)


# ---------------- simulado.xlsx ----------------

def baixar_workbook(pasta_estudo_id: str) -> Workbook:
    arquivo = _buscar("simulado.xlsx", pasta_estudo_id)
    if not arquivo:
        wb = Workbook()
        wb.remove(wb.active)
        return wb
    dados = _baixar_bytes(arquivo["id"])
    return load_workbook(io.BytesIO(dados))


def salvar_workbook(pasta_estudo_id: str, wb: Workbook):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    arquivo = _buscar("simulado.xlsx", pasta_estudo_id)
    _upload_bytes("simulado.xlsx", buffer.read(), pasta_estudo_id, MIME_XLSX,
                  file_id=arquivo["id"] if arquivo else None)
