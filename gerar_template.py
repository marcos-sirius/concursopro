"""
Gera o template Excel de CONTEÚDO PROGRAMÁTICO.
Rode uma vez: python gerar_template.py
Depois preencha o arquivo gerado para cada novo estudo/concurso.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def gerar_template(caminho="template_conteudo_programatico.xlsx"):
    wb = Workbook()

    # ── Guia "Eixos" ─────────────────────────────
    ws_eixos = wb.active
    ws_eixos.title = "Eixos"
    ws_eixos.append(["eixo", "qtd_questoes", "pontos_por_questao", "numero_inicial_questao", "tamanho_bloco"])
    ws_eixos.append(["Específicos", 50, 24, 1, 10])
    ws_eixos.append(["Língua Portuguesa", 10, 12, 51, 10])
    ws_eixos.append(["Noções de Informática", 5, 8, 61, 5])
    ws_eixos.append(["Legislação e Normas", 10, 12, 66, 10])

    # ── Guia "Temas" (conteúdo programático com peso 80/20) ──
    ws_temas = wb.create_sheet("Temas")
    ws_temas.append(["eixo", "tema", "peso"])
    exemplos = [
        ("Específicos", "Avaliação Psicológica: laudo, relatório, parecer e atestado.", 8),
        ("Específicos", "Ética profissional: sigilo profissional e código de ética.", 8),
        ("Língua Portuguesa", "Sintaxe: termos da oração, concordância e regência.", 8),
        ("Língua Portuguesa", "Mecanismos de Textualidade: coesão e coerência.", 7),
        ("Noções de Informática", "Suítes de Escritório: Word, Excel, PowerPoint.", 8),
        ("Legislação e Normas", "Regime Jurídico Único: Lei nº 8.112/1990.", 8),
    ]
    for linha in exemplos:
        ws_temas.append(linha)

    # Formatação simples de cabeçalho
    for ws in (ws_eixos, ws_temas):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        for col in ws.columns:
            max_len = max(len(str(c.value)) for c in col if c.value is not None)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    wb.save(caminho)
    print(f"Template gerado: {caminho}")

if __name__ == "__main__":
    gerar_template()
