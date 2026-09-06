"""
Lê o template_conteudo_programatico.xlsx e devolve a estrutura de eixos.
"""
from openpyxl import load_workbook


def carregar_conteudo_programatico(caminho_excel) -> dict:
    """
    Aceita tanto um caminho de arquivo (str/Path, uso local/CLI) quanto um
    objeto tipo-arquivo em memória (ex: o retorno de st.file_uploader no
    Streamlit, ou um BytesIO), para funcionar também no app web sem precisar
    tocar em disco.

    Retorna:
    {
        "Específicos": {
            "qtd_questoes": 30,
            "pontos_por_questao": 24,
            "numero_inicial_questao": 1,
            "temas": {"tema 1": 8, "tema 2": 7, ...}
        },
        ...
    }
    """
    eh_caminho = isinstance(caminho_excel, str) or hasattr(caminho_excel, "__fspath__")

    if eh_caminho:
        from pathlib import Path

        caminho = Path(caminho_excel)
        if not caminho.exists():
            raise FileNotFoundError(
                f"Arquivo não encontrado: {caminho_excel}\n"
                "Dica: cole o caminho sem aspas, ou copie o arquivo para a pasta do projeto."
            )
        if caminho.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            raise ValueError(
                f"Formato não suportado: {caminho.suffix}. Use um arquivo .xlsx."
            )
        wb = load_workbook(caminho_excel, data_only=True)
    else:
        # objeto tipo-arquivo (UploadedFile do Streamlit, BytesIO, etc.)
        wb = load_workbook(caminho_excel, data_only=True)

    if "Eixos" not in wb.sheetnames or "Temas" not in wb.sheetnames:
        raise ValueError(
            "O Excel precisa ter as guias 'Eixos' e 'Temas'. "
            "Use template_conteudo_programatico.xlsx como base."
        )

    ws_eixos = wb["Eixos"]
    eixos = {}
    for row in ws_eixos.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        nome, qtd, pontos, num_inicial = row[0], row[1], row[2], row[3]
        tamanho_bloco = row[4] if len(row) > 4 and row[4] else 10
        eixos[nome] = {
            "qtd_questoes": int(qtd),
            "pontos_por_questao": int(pontos),
            "numero_inicial_questao": int(num_inicial),
            "tamanho_bloco": int(tamanho_bloco),
            "temas": {},
        }

    ws_temas = wb["Temas"]
    for row in ws_temas.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        eixo, tema, peso = row[0], row[1], row[2]
        if eixo not in eixos:
            raise ValueError(
                f"Eixo '{eixo}' aparece na guia Temas mas não está cadastrado na guia Eixos."
            )
        eixos[eixo]["temas"][tema] = int(peso) if peso else 1

    for eixo, dados in eixos.items():
        if not dados["temas"]:
            raise ValueError(f"O eixo '{eixo}' não tem nenhum tema cadastrado na guia Temas.")

    return eixos


def importar_temas_de_docx(caminho_docx: str) -> list:
    """
    Auxiliar opcional: extrai uma lista de linhas de um Word com o edital/conteúdo
    programático, para facilitar colar na guia 'Temas' depois (sem peso automático,
    o peso 80/20 você define manualmente — isso exige julgamento sobre o edital).
    """
    from docx import Document  # python-docx

    doc = Document(caminho_docx)
    linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return linhas
