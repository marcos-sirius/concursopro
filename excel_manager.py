"""
simulado.xlsx
 ├── <Eixo 1>      -> resetada a cada simulado novo (buffer de trabalho)
 ├── <Eixo 2>      -> idem
 ├── ...
 └── Consolidado   -> nunca é limpa; cada linha tem o número da simulação na 1ª coluna
"""
from pathlib import Path
from openpyxl import Workbook, load_workbook

CABECALHO_EIXO = ["pergunta", "opcao_a", "opcao_b", "opcao_c", "opcao_d", "opcao_e", "correta", "comentario", "tema"]
CABECALHO_CONSOLIDADO = ["simulacao", "eixo"] + CABECALHO_EIXO


def abrir_ou_criar(excel_path: str) -> Workbook:
    caminho = Path(excel_path)
    if caminho.exists():
        return load_workbook(excel_path)
    wb = Workbook()
    wb.remove(wb.active)  # remove a sheet default em branco
    ws = wb.create_sheet("Consolidado")
    ws.append(CABECALHO_CONSOLIDADO)
    wb.save(excel_path)
    return wb


def resetar_guia_eixo(wb: Workbook, eixo: str):
    """Limpa (ou cria) a guia do eixo, deixando só o cabeçalho — equivalente ao clearContents()."""
    if eixo in wb.sheetnames:
        wb.remove(wb[eixo])
    ws = wb.create_sheet(eixo)
    ws.append(CABECALHO_EIXO)
    return ws


def gravar_questoes_no_eixo(wb: Workbook, eixo: str, tema: str, questoes: list):
    ws = wb[eixo] if eixo in wb.sheetnames else resetar_guia_eixo(wb, eixo)
    for q in questoes:
        opcoes = q["opcoes"]
        # garante 5 colunas mesmo se o modelo devolver 4
        while len(opcoes) < 5:
            opcoes.append("")
        ws.append([
            q["pergunta"],
            opcoes[0], opcoes[1], opcoes[2], opcoes[3], opcoes[4],
            q["correta"],
            q.get("comentario", ""),
            tema,
        ])


def consolidar(wb: Workbook, numero_simulacao: int, eixos_questoes: dict):
    """
    eixos_questoes: {eixo: [(tema, [questoes]), ...]}
    Um eixo pode ter vários blocos (um por tema sorteado), já que a
    distribuição agora mescla vários temas por eixo em vez de um só.
    Acrescenta tudo na guia Consolidado com o número da simulação na 1ª coluna.
    """
    if "Consolidado" not in wb.sheetnames:
        ws = wb.create_sheet("Consolidado")
        ws.append(CABECALHO_CONSOLIDADO)
    else:
        ws = wb["Consolidado"]

    for eixo, blocos in eixos_questoes.items():
        for tema, questoes in blocos:
            for q in questoes:
                opcoes = q["opcoes"]
                while len(opcoes) < 5:
                    opcoes.append("")
                ws.append([
                    numero_simulacao,
                    eixo,
                    q["pergunta"],
                    opcoes[0], opcoes[1], opcoes[2], opcoes[3], opcoes[4],
                    q["correta"],
                    q.get("comentario", ""),
                    tema,
                ])


def perguntas_ja_usadas(wb: Workbook, eixo: str, tema: str = None, limite: int = 30) -> list:
    """
    Lê o histórico do Consolidado para reduzir repetição no prompt da IA,
    sem deixar o prompt crescer sem limite ao longo de muitos simulados:

      - Filtra por tema (não pelo eixo inteiro): só interessa evitar repetição
        do MESMO assunto que está sendo gerado agora.
      - Pega as últimas `limite` ocorrências (mais recentes), não o histórico todo.
      - Resume cada pergunta às ~10 primeiras palavras: suficiente pra IA
        reconhecer um padrão repetido, custando uma fração dos tokens.
    """
    if "Consolidado" not in wb.sheetnames:
        return []
    ws = wb["Consolidado"]

    encontradas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        eixo_row, tema_row, pergunta = row[1], row[10], row[2]
        if eixo_row != eixo:
            continue
        if tema is not None and tema_row != tema:
            continue
        if pergunta:
            encontradas.append(pergunta)

    recentes = encontradas[-limite:]
    resumidas = [" ".join(str(p).split()[:10]) + "..." for p in recentes]
    return resumidas


def questoes_por_tema(wb: Workbook, eixo: str, tema: str) -> list:
    """
    Lê TODAS as questões já geradas pra esse (eixo, tema) — em qualquer
    simulado desse estudo — direto do Consolidado. Usado só na hora de
    REAPROVEITAR conteúdo pra outro estudo (banco_compartilhado_manager.py);
    não faz parte do fluxo normal de geração.

    Devolve no formato {"pergunta", "opcoes", "correta" (índice 0-4),
    "comentario"} — o mesmo formato que gravar_questoes_no_eixo espera.
    Não repete a mesma pergunta duas vezes, mesmo que ela apareça em vários
    simulados antigos desse estudo.
    """
    if "Consolidado" not in wb.sheetnames:
        return []
    ws = wb["Consolidado"]
    letras = ["A", "B", "C", "D", "E"]

    encontradas = []
    vistas = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        eixo_row = row[1]
        tema_row = row[10] if len(row) > 10 else None
        if eixo_row != eixo or tema_row != tema:
            continue

        pergunta = row[2]
        if not pergunta or pergunta in vistas:
            continue
        vistas.add(pergunta)

        opcoes = [row[3], row[4], row[5], row[6], row[7]]
        correta_valor = row[8]
        comentario = row[9] if len(row) > 9 and row[9] is not None else ""

        if isinstance(correta_valor, str) and correta_valor.strip().upper() in letras:
            idx_correta = letras.index(correta_valor.strip().upper())
        else:
            idx_correta = int(correta_valor) if correta_valor is not None else 0

        encontradas.append({
            "pergunta": pergunta,
            "opcoes": opcoes,
            "correta": idx_correta,
            "comentario": comentario,
        })
    return encontradas
